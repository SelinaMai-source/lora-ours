#!/usr/bin/env python3
"""Generate 实验v3.docx — comprehensive Chinese experiment report (CCF-A style)."""

from __future__ import annotations

import csv
import json
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from docx.oxml import OxmlElement

PROJECT_ROOT = Path(__file__).resolve().parents[1]
DOC_DIR = PROJECT_ROOT / "文档记录"
CHART_DIR = DOC_DIR / "_charts_v3"
REPORT_DATE = datetime.now().strftime("%Y-%m-%d")

# ---------------------------------------------------------------------------
# Metrics paths
# ---------------------------------------------------------------------------
METRIC_FILES = {
    "Ours (v4_sota_39)": PROJECT_ROOT / "results/runs/paper_instrdialog_ours_full_s123_v4_sota_39/final_metrics.json",
    "Ours (v6_sota_2)": PROJECT_ROOT / "results/runs/paper_instrdialog_ours_full_s123_v6_sota_2/final_metrics.json",
    "Ours (v7_sota_1, 进行中)": PROJECT_ROOT / "results/runs/paper_instrdialog_ours_full_s123_v7_sota_1/final_metrics.json",
    "Sequential (true_fair)": PROJECT_ROOT / "results/runs/paper_instrdialog_seq_s123_true_fair/final_metrics.json",
    "Replay(10) (旧协议 s123)": PROJECT_ROOT / "results/runs/paper_instrdialog_replay_b10_s123/final_metrics.json",
    "Replay(50) (true_fair)": PROJECT_ROOT / "results/runs/paper_instrdialog_replay_b50_s123_true_fair/final_metrics.json",
    "PeriodicLatest (true_fair)": PROJECT_ROOT / "results/runs/paper_instrdialog_periodic_latest_s123_true_fair/final_metrics.json",
    "BankNoRouter (true_fair)": PROJECT_ROOT / "results/runs/paper_instrdialog_bank_no_router_s123_true_fair/final_metrics.json",
    "RouterOnly (true_fair)": PROJECT_ROOT / "results/runs/paper_instrdialog_router_only_s123_true_fair/final_metrics.json",
}

SEGMENT_CSV = {
    "v4_sota_39": PROJECT_ROOT / "results/tables/paper_instrdialog_ours_full_s123_v4_sota_39_segment_metrics.csv",
    "v6_sota_2": PROJECT_ROOT / "results/tables/paper_instrdialog_ours_full_s123_v6_sota_2_segment_metrics.csv",
}

SOTA_TARGETS = {
    "eval.seen_avg_score": (">=", 0.600),
    "eval.seen_avg_task_aware_score": (">=", 0.600),
    "eval.forgetting": ("<=", 0.0333),
    "eval.task_aware_forgetting": ("<=", 0.077),
    "eval.token_f1_mean": (">=", 0.600),
    "eval.lcs_overlap_mean": (">=", 0.600),
}

MAIN_METRIC_KEYS = [
    ("eval.seen_avg_score", "Seen Avg Score"),
    ("eval.seen_avg_task_aware_score", "Seen Avg Task-Aware"),
    ("eval.forgetting", "Forgetting"),
    ("eval.task_aware_forgetting", "Task-Aware Forgetting"),
    ("eval.token_f1_mean", "Token F1 Mean"),
    ("eval.lcs_overlap_mean", "LCS Overlap Mean"),
    ("routing.oracle_agreement_rate", "Oracle Agreement"),
    ("drift.false_alarm_rate", "Drift False Alarm"),
    ("drift.miss_rate", "Drift Miss Rate"),
    ("drift.detection_delay_mean", "Drift Delay (seg)"),
]

GAPS: List[str] = []


def resolve_output_path() -> Path:
    """If 实验v2.docx exists use v3; if v3 exists use next number."""
    if not (DOC_DIR / "实验v2.docx").exists():
        # No v2 docx — user requested v3
        candidate = DOC_DIR / "实验v3.docx"
        if candidate.exists():
            n = 4
            while (DOC_DIR / f"实验v{n}.docx").exists():
                n += 1
            return DOC_DIR / f"实验v{n}.docx"
        return candidate
    n = 3
    while (DOC_DIR / f"实验v{n}.docx").exists():
        n += 1
    return DOC_DIR / f"实验v{n}.docx"


def load_metrics(path: Path) -> Optional[Dict[str, Any]]:
    if not path.is_file():
        GAPS.append(f"缺失指标文件: {path.relative_to(PROJECT_ROOT)}")
        return None
    with path.open(encoding="utf-8") as f:
        data = json.load(f)
    final = dict(data.get("final", {}))
    drift_q = data.get("drift_quality") or {}
    for k, v in drift_q.items():
        final[f"drift.{k.replace('num_', '')}"] = v
        if k == "false_alarm_rate":
            final["drift.false_alarm_rate"] = v
        elif k == "miss_rate":
            final["drift.miss_rate"] = v
        elif k == "detection_delay_mean":
            final["drift.detection_delay_mean"] = v
    return final


def load_v7_partial() -> Optional[Dict[str, Any]]:
    jsonl = PROJECT_ROOT / "results/runs/paper_instrdialog_ours_full_s123_v7_sota_1/metrics.jsonl"
    if not jsonl.is_file():
        GAPS.append("v7_sota_1 尚无 metrics.jsonl（运行可能尚未启动）")
        return None
    last = None
    with jsonl.open(encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if line:
                last = json.loads(line)
    if last is None:
        GAPS.append("v7_sota_1 metrics.jsonl 为空")
        return None
    seg = last.get("segment_id", "?")
    GAPS.append(f"v7_sota_1 尚未完成：使用 seg{seg} 部分指标（非 final_metrics.json）")
    return last


def fmt(v: Any, digits: int = 3) -> str:
    if v is None:
        return "—"
    if isinstance(v, float):
        if v != v:  # NaN
            return "—"
        return f"{v:.{digits}f}"
    return str(v)


def set_doc_font(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    for level in range(1, 4):
        hs = doc.styles[f"Heading {level}"]
        hs.font.name = "Times New Roman"
        hs._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        hs.font.color.rgb = RGBColor(0, 0, 0)


def add_title(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(18)
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_para(doc: Document, text: str, bold: bool = False) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold


def add_bullet(doc: Document, text: str) -> None:
    doc.add_paragraph(text, style="List Bullet")


def add_code_block(doc: Document, text: str) -> None:
    for line in text.strip().splitlines():
        p = doc.add_paragraph()
        run = p.add_run(line)
        run.font.name = "Courier New"
        run.font.size = Pt(9)
        p.paragraph_format.left_indent = Cm(0.8)


def add_table(doc: Document, headers: List[str], rows: List[List[str]], highlight_col: Optional[int] = None) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        for p in hdr[i].paragraphs:
            for r in p.runs:
                r.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = val
            if highlight_col is not None and i == highlight_col:
                for p in cells[i].paragraphs:
                    for r in p.runs:
                        r.bold = True
    doc.add_paragraph()


def read_segment_trajectory(csv_path: Path) -> Tuple[List[int], List[float], List[float]]:
    segs, seen, oracle = [], [], []
    with csv_path.open(encoding="utf-8") as f:
        reader = csv.DictReader(f)
        for row in reader:
            segs.append(int(row["segment_id"]))
            seen.append(float(row["eval.seen_avg_score"]))
            oracle.append(float(row["routing.oracle_agreement_rate"]))
    return segs, seen, oracle


def make_charts(all_metrics: Dict[str, Dict[str, Any]]) -> Dict[str, Path]:
    CHART_DIR.mkdir(parents=True, exist_ok=True)
    paths: Dict[str, Path] = {}

    # --- Chart 1: main metrics bar comparison ---
    methods = [
        "Sequential (true_fair)",
        "Replay(50) (true_fair)",
        "PeriodicLatest (true_fair)",
        "BankNoRouter (true_fair)",
        "RouterOnly (true_fair)",
        "Ours (v4_sota_39)",
        "Ours (v6_sota_2)",
    ]
    metric_key = "eval.seen_avg_task_aware_score"
    labels = [m.replace(" (true_fair)", "").replace("Ours ", "Ours\n") for m in methods]
    values = []
    for m in methods:
        d = all_metrics.get(m, {})
        values.append(float(d.get(metric_key, 0) or 0))

    fig, ax = plt.subplots(figsize=(10, 5))
    colors = ["#95a5a6"] * 5 + ["#3498db", "#e74c3c"]
    bars = ax.bar(range(len(methods)), values, color=colors, edgecolor="black", linewidth=0.5)
    ax.axhline(0.600, color="#27ae60", linestyle="--", linewidth=1.5, label="SOTA 目标 (0.600)")
    ax.set_xticks(range(len(methods)))
    ax.set_xticklabels(labels, rotation=25, ha="right", fontsize=9)
    ax.set_ylabel("Seen Avg Task-Aware Score")
    ax.set_title("InstrDialog s123：Ours vs True-Fair Baselines（终局指标）")
    ax.set_ylim(0, max(0.65, max(values) * 1.15))
    ax.legend(loc="upper right")
    for bar, v in zip(bars, values):
        ax.text(bar.get_x() + bar.get_width() / 2, bar.get_height() + 0.01, f"{v:.3f}", ha="center", va="bottom", fontsize=8)
    fig.tight_layout()
    p1 = CHART_DIR / "main_task_aware_comparison.png"
    fig.savefig(p1, dpi=150)
    plt.close(fig)
    paths["main"] = p1

    # --- Chart 2: six SOTA targets radar-like grouped bar ---
    ours_v6 = all_metrics.get("Ours (v6_sota_2)", {})
    target_names = ["Seen\nAvg", "Task-Aware\nSeen", "Token F1", "LCS\nOverlap"]
    ours_vals = [
        float(ours_v6.get("eval.seen_avg_score", 0) or 0),
        float(ours_v6.get("eval.seen_avg_task_aware_score", 0) or 0),
        float(ours_v6.get("eval.token_f1_mean", 0) or 0),
        float(ours_v6.get("eval.lcs_overlap_mean", 0) or 0),
    ]
    targets = [0.600, 0.600, 0.600, 0.600]
    x = range(len(target_names))
    w = 0.35
    fig, ax = plt.subplots(figsize=(8, 5))
    ax.bar([i - w / 2 for i in x], ours_vals, w, label="Ours v6_sota_2", color="#e74c3c")
    ax.bar([i + w / 2 for i in x], targets, w, label="SOTA 目标", color="#27ae60", alpha=0.7)
    ax.set_xticks(list(x))
    ax.set_xticklabels(target_names)
    ax.set_ylabel("Score")
    ax.set_title("六项 SOTA 目标：v6_sota_2 当前 vs 目标（越高越好子集）")
    ax.legend()
    ax.set_ylim(0, 0.75)
    fig.tight_layout()
    p2 = CHART_DIR / "sota_gap_chart.png"
    fig.savefig(p2, dpi=150)
    plt.close(fig)
    paths["sota_gap"] = p2

    # --- Chart 3: segment trajectory ---
    fig, ax1 = plt.subplots(figsize=(10, 5))
    for label, key in [("v4_sota_39", "v4_sota_39"), ("v6_sota_2", "v6_sota_2")]:
        csv_p = SEGMENT_CSV.get(key)
        if csv_p and csv_p.is_file():
            segs, seen, _ = read_segment_trajectory(csv_p)
            ax1.plot(segs, seen, marker="o", markersize=4, label=f"seen_avg ({label})")
    ax1.axhline(0.600, color="#27ae60", linestyle="--", alpha=0.7, label="SOTA 目标")
    ax1.set_xlabel("Segment ID")
    ax1.set_ylabel("eval.seen_avg_score")
    ax1.set_title("Ours 训练轨迹：seen_avg_score 随 segment 变化")
    ax1.legend(loc="lower right", fontsize=9)
    ax1.grid(True, alpha=0.3)
    fig.tight_layout()
    p3 = CHART_DIR / "trajectory_seen_avg.png"
    fig.savefig(p3, dpi=150)
    plt.close(fig)
    paths["trajectory"] = p3

    # --- Chart 4: oracle agreement trajectory ---
    fig, ax = plt.subplots(figsize=(10, 5))
    for label, key, color in [
        ("v4_sota_39 (legacy router)", "v4_sota_39", "#3498db"),
        ("v6_sota_2 (prototype)", "v6_sota_2", "#e74c3c"),
    ]:
        csv_p = SEGMENT_CSV.get(key)
        if csv_p and csv_p.is_file():
            segs, _, oracle = read_segment_trajectory(csv_p)
            ax.plot(segs, oracle, marker="s", markersize=4, label=label, color=color)
    ax.axhline(0.45, color="#f39c12", linestyle=":", label="v7 判据 (0.45)")
    ax.set_xlabel("Segment ID")
    ax.set_ylabel("routing.oracle_agreement_rate")
    ax.set_title("路由质量轨迹：Oracle Agreement Rate")
    ax.legend(fontsize=9)
    ax.grid(True, alpha=0.3)
    fig.tight_layout()
    p4 = CHART_DIR / "trajectory_oracle.png"
    fig.savefig(p4, dpi=150)
    plt.close(fig)
    paths["oracle"] = p4

    return paths


def build_document(out_path: Path) -> None:
    all_metrics: Dict[str, Dict[str, Any]] = {}
    for name, path in METRIC_FILES.items():
        if name.startswith("Ours (v7"):
            if path.is_file():
                all_metrics[name] = load_metrics(path) or {}
            else:
                partial = load_v7_partial()
                all_metrics[name] = partial or {}
            continue
        m = load_metrics(path)
        if m:
            all_metrics[name] = m

    chart_paths = make_charts(all_metrics)

    doc = Document()
    set_doc_font(doc)

    add_title(doc, "持续指令微调实验报告（InstrDialog · Ours SOTA 路线）")
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run(f"版本：实验报告 v3  |  生成日期：{REPORT_DATE}  |  项目：Lora-code").italic = True
    doc.add_paragraph()

    # ===== 0) 实验规则 =====
    add_heading(doc, "0  实验规则与工程规范", level=1)
    add_para(doc, "为保证长周期 GPU 实验的可复现性与可恢复性，所有正式 SOTA 跑分必须遵守以下三条硬性规则：")
    add_bullet(doc, "必须挂 tmux：所有训练/评估均在 tmux 会话中后台运行（如 v6_sota_2、v7_sota_1），禁止依赖易断开的交互式 shell。")
    add_bullet(doc, "必须设置自动断点续传：output.save_every_segment=true，每 segment 持久化 adapter checkpoint、metrics.jsonl、branch_registry.json；异常中断后可从最近 segment 恢复。")
    add_bullet(doc, "必须在 Weights & Biases 上清晰有框架地记录：project=lora-citb-sota，按版本分组（group=v6_sota / v7_sota），tags 含版本号与方法组件；同步 train/eval/routing/drift 全量标量。")
    add_para(doc, "典型 tmux 启动示例：")
    add_code_block(doc, """tmux new-session -d -s v7_sota_1 'bash archive/v7_sota_1/run_v7_sota_1.sh 2>&1 | tee v7_sota_1.log'
tmux attach -t v7_sota_1   # 附着查看，Ctrl+B D 分离""")
    add_para(doc, "W&B 记录框架：")
    add_table(
        doc,
        ["字段", "Ours SOTA 示例", "Baseline 示例"],
        [
            ["project", "lora-citb-sota", "lora-citb-acl"],
            ["group", "v7_sota", "paper_baselines_20260530"],
            ["tags", "v7_sota_1, prototype_router, verify_then_route", "seq, true_fair, instrdialog"],
            ["notes", "单点假设 + 轨迹早停判据", "方法变体 + 公平协议说明"],
        ],
    )

    # ===== 1) Contribution Points =====
    add_heading(doc, "1  目前最新 Ours 的贡献点（v6 → v7）", level=1)
    add_para(
        doc,
        "本方法面向 task-agnostic 持续指令微调流：在 Llama-3.1-8B-Instruct 上，通过漂移检测触发 LoRA 分支增殖、"
        "漂移锚定原型路由器分配样本、anti-overlap 抑制分支参数重叠，并在 v7 引入 verify-then-route 不确定性门控。"
        "以下贡献点均可在代码库 archive/v6_sota_2/CHANGES.md、archive/v7_sota_1/CHANGES.md 及 core/ 中溯源。",
    )

    add_heading(doc, "1.1  双层锚点漂移检测（Drift Detector）", level=2)
    add_para(doc, "核心思想：在固定锚点集上监测 teacher-forced answer NLL，用校准 CUSUM 检测分布偏移，触发新 LoRA 分支 spawn。")
    add_para(doc, "锚点构建：K-Center Greedy 或分 segment 轮询采样，划分为 core（易样本）与 probe（难样本）两层：")
    add_para(doc, "  • core/probe 划分遵循 curriculum_strategy（默认 easy_core_hard_probe）")
    add_para(doc, "  • anchor_source=train（v6+ 公平性修复：不使用 eval 集作监控，避免标签泄漏）")
    add_para(doc, "  • force_spawn_on_segment_boundary=false（禁止利用任务边界 oracle 信息）")
    add_para(doc, "CUSUM 更新（probe 侧，简化记法）：")
    add_para(doc, "  S_probe ← max(0, S_probe + (NLL_probe − baseline_probe) − slack_probe)")
    add_para(doc, "  触发条件：S_probe ≥ τ_probe 且连续命中次数 ≥ min_consecutive_probe_hits")
    add_para(doc, "公平性修复摘要：v6_sota_1 起统一 train-split 锚点、关闭 segment 边界强制 spawn、"
               "关闭 meta_threshold 对任务边界的隐式泄漏；与 true_fair baseline 协议对齐。")

    add_heading(doc, "1.2  漂移锚定原型路由器（Drift-Anchored Prototype Router, v6_sota_2）", level=2)
    add_para(doc, "动机：v6_sota_1 的 legacy router（Poincaré 双曲距离 + PID + Tsallis + Ising 堆叠）导致路由概率近均匀，"
               "oracle_agreement 从 seg4 的 0.49 跌至 seg12 的 0.15。")
    add_para(doc, "原型构建：每 segment 训练后，用 teacher-forced NLL 自一致伪标签将样本指派给最优分支，"
               "对被指派样本的 L2 归一化隐层特征取均值，EMA 更新分支原型：")
    add_para(doc, "  p_b ← normalize( α · p_b + (1−α) · mean({f(x_i) | label_i = b}) ),  α = prototype_ema = 0.8")
    add_para(doc, "冻结一致性：bank 冻结分支 b 时，其原型同步冻结（仅允许一次性初始化），保证「路由记忆」与「参数记忆」一致。")
    add_para(doc, "打分函数（余弦相似度 + temperature softmax，支持 soft top-3 混合）：")
    add_para(doc, "  sim_b(x) = cos(f(x), p_b) = f̂(x) · p̂_b")
    add_para(doc, "  P(b|x) = softmax(sim_b(x) / T)")
    add_para(doc, "伪代码（原型更新与路由）：")
    add_code_block(
        doc,
        """# 训练后：伪标签 + 原型 EMA
for (x, feat) in segment_train:
    b* = argmin_b NLL(x; adapter_b)   # teacher-forced 自一致
    assign pseudo_label = b*
for branch b in branches:
    if b frozen and prototype exists: continue
    proto[b] = EMA(proto[b], mean(feat[pseudo_label==b]))

# 推理：余弦路由 + soft top-3
scores = softmax(cos(f(x), proto) / T)
route = top_k_mix(scores, k=3)  # 参数空间加权混合""",
    )

    add_heading(doc, "1.3  置信度校准 + Verify-then-Route（v7_sota_1）", level=2)
    add_para(doc, "假设：原型路由误差来自 (1) 低支持度原型的过度自信；(2) 低 margin 时 soft blending 的破坏性干扰。")
    add_para(doc, "支持度校准（贝叶斯式可靠性加权）：")
    add_para(doc, "  sim_cal_b = sim_b · count_b / (count_b + prior),  prior = 10")
    add_para(doc, "Verify-then-Route：当 top1−top2 概率 margin < arbitration_margin (=0.12) 时，"
               "对 top-k 候选分支计算 label-free prompt NLL，硬路由到 NLL 最小者，跳过 soft blending：")
    add_code_block(
        doc,
        """margin = P_top1 - P_top2
if margin < 0.12:
    candidates = top_k(probs, k=3)
    b* = argmin_{b in candidates} NLL_prompt(x; adapter_b)
    hard_route(b*)   # 不做参数混合
else:
    soft_blend(top_k=3)  # 保持 v6_sota_2 高置信收益""",
    )
    add_para(doc, "训练预算保持不变：epochs=1, batch_size=2（v6_3~6 实验证明 3ep/5ep 或纯硬路由会破坏专精分支）。")

    add_heading(doc, "1.4  Anti-Overlap 与分支银行", level=2)
    add_para(doc, "Anti-overlap 损失：对多分支激活表示与 LoRA 权重施加正交/低余弦惩罚，β=0.05（activation:weight = 6:4）。")
    add_para(doc, "LoRA Bank：max_branches=15，spawn_on_drift=true，freeze_old_branches=true；"
               "新 segment 仅训练 active 分支，旧分支参数冻结。")

    add_heading(doc, "1.5  与相近工作的区别（Novelty）", level=2)
    add_table(
        doc,
        ["相近工作", "本方法区别"],
        [
            ["S-Prompts / HiDe-Prompt", "任务边界未知；原型由 NLL 自一致验证而非 K-Means 聚类"],
            ["MoE-Adapters / DEMix", "校准对象为非监督 CL 动态 spawn 原型；NLL 仲裁与 soft blend 按 margin 解耦"],
            ["O-LoRA", "子空间正交但无路由；本方法显式维护路由-参数冻结一致性"],
            ["Temperature Scaling (Guo et al.)", "校准支持度来自伪标签计数，而非分类头温度"],
        ],
    )

    # ===== 2) Baselines =====
    add_heading(doc, "2  目前的 Baseline 定义与公平协议", level=1)
    add_para(
        doc,
        "主表 baseline 采用 true_fair 协议（epochs_per_segment=5, batch_size=1, lr=2e-4），"
        "与 instrdialog 19 segments、seed=123 对齐。配置文件位于 configs/paper/instrdialog__*__s123.yaml；"
        "true_fair 跑分 run_name 后缀 _true_fair。",
    )
    add_table(
        doc,
        ["方法", "variant_id", "核心机制", "true_fair 协议"],
        [
            ["Sequential", "seq", "单 LoRA 顺序微调，无 bank/router", "5 ep, batch=1"],
            ["Replay(10)", "replay_b10", "buffer=10, replay_ratio=0.3", "5 ep, batch=1（true_fair 待补）"],
            ["Replay(50)", "replay_b50", "buffer=50", "5 ep, batch=1 ✓"],
            ["PeriodicLatest", "periodic_latest", "每 2 segment spawn，始终用最新分支", "5 ep, batch=1 ✓"],
            ["BankNoRouter", "bank_no_router", "有 drift+bank，无 router", "5 ep, batch=1 ✓"],
            ["RouterOnly", "router_only", "仅 router 选分支，无 drift 触发", "5 ep, batch=1 ✓"],
        ],
    )
    add_para(doc, "Ours（SOTA 线）刻意使用更低训练预算（1 ep, batch=2）以控制算力；"
               "v6_sota_3 曾尝试与 baseline 对齐 5ep 但导致 soft blend 破坏专精分支。")

    # ===== 3) 实验设计 =====
    add_heading(doc, "3  实验设计", level=1)
    add_table(
        doc,
        ["维度", "设定"],
        [
            ["骨干模型", "Llama-3.1-8B-Instruct (bf16, max_seq_len=512)"],
            ["PEFT", "LoRA r=16, α=32, target=[q_proj, v_proj]"],
            ["Benchmark", "InstrDialog：19 segments, seed=123, 50 train + 10 eval / segment"],
            ["生成", "greedy, max_new_tokens=64, mask_eos_token_in_labels=true"],
            ["方法流水线", "漂移检测 → 分支 spawn → 原型路由(+NLL仲裁) → segment 级 eval"],
            ["学习范式", "Segment-by-segment continual learning，每段结束评估全部已见任务"],
            ["当前运行", "v7_sota_1（tmux: v7_sota_1，config: configs/paper/v7_sota_1.yaml）"],
        ],
    )
    add_para(doc, "流程示意（Mermaid 逻辑）：")
    add_code_block(
        doc,
        """for seg in stream[0..18]:
    train LoRA on seg.train (routed via prototype router)
    update prototypes with NLL pseudo-labels
    monitor anchor NLL → maybe spawn new branch + freeze old
    eval on seg + all seen segments → log metrics + wandb""",
    )

    # ===== 4) Expected Output =====
    add_heading(doc, "4  期望输出与 SOTA 指标", level=1)
    add_para(doc, "终局（segment 18）六项 SOTA 目标需同时满足：")
    rows_sota = []
    v6 = all_metrics.get("Ours (v6_sota_2)", {})
    for key, (op, target) in SOTA_TARGETS.items():
        cur = v6.get(key)
        cur_f = float(cur) if cur is not None else None
        if cur_f is None:
            status = "缺失"
        elif op == ">=":
            status = "✓ 达标" if cur_f >= target else f"未达标 (差距 {target - cur_f:+.3f})"
        else:
            status = "✓ 达标" if cur_f <= target else f"未达标 (超出 {cur_f - target:+.3f})"
        rows_sota.append([key, f"{op} {target:.4f}", fmt(cur_f), status])
    add_table(doc, ["指标", "SOTA 目标", "v6_sota_2 当前", "状态"], rows_sota)

    add_para(doc, "附加路由/漂移质量指标：")
    add_bullet(doc, "routing.oracle_agreement_rate：路由选中分支 vs NLL-oracle 最佳分支的一致率（v6≈0.37，v4≈0.49）")
    add_bullet(doc, "routing.decision_confidence_mean / decision_entropy_mean：决策置信度与熵")
    add_bullet(doc, "drift.false_alarm_rate / miss_rate / detection_delay_mean：漂移检测质量（当前 miss_rate≈0.5）")
    add_bullet(doc, "overlap_mean_cosine：分支间 LoRA 权重余弦（越低越好，v6≈0.88）")
    add_bullet(doc, "v7 特有：nll_arbitration_changed / num_routed 比率（期望 5–30%）")

    # ===== 5) Results =====
    add_heading(doc, "5  最佳 Baselines 与 Ours 结果对比", level=1)

    # Main comparison table
    compare_methods = [
        "Sequential (true_fair)",
        "Replay(10) (旧协议 s123)",
        "Replay(50) (true_fair)",
        "PeriodicLatest (true_fair)",
        "BankNoRouter (true_fair)",
        "RouterOnly (true_fair)",
        "Ours (v4_sota_39)",
        "Ours (v6_sota_2)",
    ]
    headers = ["方法"] + [label for _, label in MAIN_METRIC_KEYS]
    rows = []
    for m in compare_methods:
        d = all_metrics.get(m, {})
        row = [m.split(" (")[0]]
        for key, _ in MAIN_METRIC_KEYS:
            if key.startswith("drift."):
                row.append(fmt(d.get(key)))
            else:
                row.append(fmt(d.get(key)))
        rows.append(row)
    add_para(doc, "表 1：终局主指标对比（InstrDialog, seed=123）")
    add_table(doc, headers, rows)

    # v7 partial if available
    v7 = all_metrics.get("Ours (v7_sota_1, 进行中)", {})
    if v7:
        add_para(doc, f"表 2：v7_sota_1 进行中部分结果（segment_id={v7.get('segment_id', '?')}，非终局）")
        add_table(
            doc,
            ["指标", "v7 部分值", "v6_sota_2 @seg18", "备注"],
            [
                ["eval.seen_avg_score", fmt(v7.get("eval.seen_avg_score")), fmt(v6.get("eval.seen_avg_score")), "轨迹对比"],
                ["eval.seen_avg_task_aware_score", fmt(v7.get("eval.seen_avg_task_aware_score")), fmt(v6.get("eval.seen_avg_task_aware_score")), ""],
                ["routing.oracle_agreement_rate", fmt(v7.get("routing.oracle_agreement_rate")), fmt(v6.get("routing.oracle_agreement_rate")), "v7 seg3≈0.61"],
                ["eval.forgetting", fmt(v7.get("eval.forgetting")), fmt(v6.get("eval.forgetting")), ""],
            ],
        )

    # Charts
    add_heading(doc, "5.1  可视化对比", level=2)
    for caption, key in [
        ("图 1：Seen Avg Task-Aware Score 终局对比", "main"),
        ("图 2：SOTA 目标差距（v6_sota_2）", "sota_gap"),
        ("图 3：seen_avg_score 训练轨迹", "trajectory"),
        ("图 4：Oracle Agreement 训练轨迹", "oracle"),
    ]:
        p = chart_paths.get(key)
        if p and p.is_file():
            doc.add_paragraph(caption)
            doc.add_picture(str(p), width=Inches(6.2))
            doc.add_paragraph()

    # Key findings
    add_heading(doc, "5.2  关键结论", level=2)
    add_bullet(doc, "Ours（v4/v6）在 seen_avg_task_aware（≈0.38–0.39）上显著优于 true_fair Sequential（0.23）与 Replay(50)（0.11），但距 SOTA 目标 0.60 仍有 ~0.21 差距。")
    add_bullet(doc, "PeriodicLatest true_fair 的 task-aware seen（0.33）最接近 Ours，但 forgetting（0.40）极高，不可接受。")
    add_bullet(doc, "RouterOnly true_fair oracle_agreement（0.53）高于 Ours v6（0.37），但 task-aware seen（0.23）低，说明无 drift 的静态路由不足以维持 CL 性能。")
    add_bullet(doc, "v6 原型路由修复了 v6_1 的路由崩溃，但 oracle 仍低于 v4 legacy router（0.49 vs 0.37）；v7 早期 seg3 oracle≈0.61，待全轨迹验证。")
    add_bullet(doc, "漂移检测 miss_rate=50% 是主要瓶颈之一：9/18 真漂移未检出，delay≈5 segments。")

    # Data gaps
    add_heading(doc, "6  数据缺口说明", level=1)
    if GAPS:
        for g in GAPS:
            add_bullet(doc, g)
    else:
        add_para(doc, "所有引用指标文件均已找到。")
    add_bullet(doc, "Replay(10) true_fair 跑分尚未完成；表中 Replay(10) 使用旧协议（5ep 前 paper matrix）数据。")
    add_bullet(doc, "v7_sota_1 运行中，final_metrics.json 尚未生成。")
    add_bullet(doc, "paper_main_results_agg.csv 中 Ours 条目仍为早期 paper_instrdialog_ours_full_s123（全零 EM），未反映 SOTA 线最新结果。")

    add_para(doc, f"— 报告由 scripts/generate_experiment_v3_docx.py 自动生成于 {REPORT_DATE} —")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))


def main() -> None:
    out = resolve_output_path()
    build_document(out)
    print(f"DOCX written: {out}")
    print(f"Size: {out.stat().st_size} bytes")
    if GAPS:
        print("Gaps:")
        for g in GAPS:
            print(f"  - {g}")


if __name__ == "__main__":
    main()
